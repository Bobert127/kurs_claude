# -*- coding: utf-8 -*-
"""Helpery do edycji plików .docx przez celowane podmiany w XML.

Dlaczego nie python-docx: szablony CV trzymają tekst w polach tekstowych
(w:txbxContent) i na wykresach (chart*.xml), których python-docx nie przechodzi.
Bezpieczniej operować na surowym XML z zachowaniem pPr/rPr i przepakować zip.

Typowe użycie:

    from docx_edit import load, get_xml, set_xml, save, rebuild_para, run, add_chart_bar, validate

    names, entries = load(CV)
    doc = get_xml(entries)                      # word/document.xml jako str
    doc = rebuild_para(doc, 'Kandydowac',       # podmień runy akapitu (zostaw pPr)
                       run('Aplikacja na stanowisko: ...', bold=True))
    set_xml(entries, doc)
    ch = get_xml(entries, 'word/charts/chart2.xml')
    ch = add_chart_bar(ch, 'n8n (automatyzacja)', 8)
    set_xml(entries, ch, 'word/charts/chart2.xml')
    save(CV, names, entries)
    print(validate(CV))
"""
import re
import zipfile
import xml.sax.saxutils as _sx


# ---------- wczytywanie / zapis ----------

def load(path):
    """Zwraca (names, entries): kolejność wpisów i mapę {nazwa: bytes}."""
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        entries = {n: z.read(n) for n in names}
    return names, entries


def get_xml(entries, name='word/document.xml'):
    return entries[name].decode('utf-8')


def set_xml(entries, text, name='word/document.xml'):
    entries[name] = text.encode('utf-8')


def save(path, names, entries):
    """Przepakowuje docx zachowując kolejność wpisów i kompresję deflate."""
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, entries[n])


# ---------- budowanie fragmentów ----------

def run(text, sz=None, bold=False, font='Candara'):
    """Pojedynczy run <w:r> z tekstem. sz w half-points (np. 18 = 9pt)."""
    rpr = '<w:rPr>'
    if font:
        rpr += '<w:rFonts w:ascii="%s" w:hAnsi="%s"/>' % (font, font)
    if bold:
        rpr += '<w:b/>'
    if sz:
        rpr += '<w:sz w:val="%s"/><w:szCs w:val="%s"/>' % (sz, sz)
    rpr += '</w:rPr>'
    return '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r>' % (rpr, _sx.escape(text))


def list_item(text, ppr, sz=18, para_id='00000001', font='Candara'):
    """Nowy element listy z podanym pPr (skopiuj pPr z istniejącego bulletu)."""
    return ('<w:p w14:paraId="%s" w14:textId="%s" w:rsidR="00000000" '
            'w:rsidRDefault="00000000">%s%s</w:p>'
            % (para_id, para_id, ppr, run(text, sz=sz, font=font)))


# ---------- edycja akapitów ----------

def _para_bounds(doc, idx):
    s = max(doc.rfind('<w:p ', 0, idx), doc.rfind('<w:p>', 0, idx))
    e = doc.index('</w:p>', idx) + len('</w:p>')
    return s, e


def rebuild_para(doc, anchor, runs_xml):
    """Zastępuje RUNY akapitu zawierającego `anchor`, zachowując jego <w:pPr>.

    Najbezpieczniejszy sposób przepisania akapitu rozbitego na wiele runów
    (np. przez znaczniki sprawdzania pisowni)."""
    j = doc.index(anchor)
    ps, pe = _para_bounds(doc, j)
    block = doc[ps:pe]
    if '</w:pPr>' in block:
        head = block[:block.index('</w:pPr>') + len('</w:pPr>')]
    else:
        head = block[:block.index('>') + 1]
    return doc[:ps] + head + runs_xml + '</w:p>' + doc[pe:]


def insert_after_para(doc, anchor, para_xml):
    """Wstawia gotowy <w:p>...</w:p> tuż po akapicie zawierającym `anchor`."""
    j = doc.index(anchor)
    _, pe = _para_bounds(doc, j)
    return doc[:pe] + para_xml + doc[pe:]


def get_ppr(doc, anchor):
    """Zwraca <w:pPr>...</w:pPr> akapitu z `anchor` (do klonowania stylu listy)."""
    j = doc.index(anchor)
    ps, pe = _para_bounds(doc, j)
    block = doc[ps:pe]
    a = block.index('<w:pPr>')
    b = block.index('</w:pPr>') + len('</w:pPr>')
    return block[a:b]


# ---------- wykresy słupkowe ----------

def _bump_cache(block, new_val, is_num):
    m = re.search(r'<c:ptCount val="(\d+)"/>', block)
    n = int(m.group(1))
    block = block[:m.start()] + '<c:ptCount val="%d"/>' % (n + 1) + block[m.end():]
    close = '</c:numCache>' if is_num else '</c:strCache>'
    pt = '<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (n, new_val)
    i = block.rindex(close)
    block = block[:i] + pt + block[i:]
    # rozszerz zakres, np. $A$2:$A$8 -> $A$2:$A$9
    block = re.sub(r'(\$[A-Z]+\$\d+:\$[A-Z]+\$)(\d+)',
                   lambda mm: mm.group(1) + str(int(mm.group(2)) + 1), block)
    return block


def add_chart_bar(chart_xml, label, value):
    """Dodaje jeden słupek do prostego wykresu słupkowego (jedna seria):
    aktualizuje kategorię (<c:cat>), wartość (<c:val>), ptCount i zakresy.
    Wartości cache wystarczają do wyświetlenia w Wordzie."""
    x = chart_xml
    cs = x.index('<c:cat>'); ce = x.index('</c:cat>') + len('</c:cat>')
    x = x[:cs] + _bump_cache(x[cs:ce], _sx.escape(str(label)), is_num=False) + x[ce:]
    vs = x.index('<c:val>'); ve = x.index('</c:val>') + len('</c:val>')
    x = x[:vs] + _bump_cache(x[vs:ve], str(value), is_num=True) + x[ve:]
    return x


# ---------- walidacja / ekstrakcja ----------

def validate(path):
    """Sprawdza, że każdy XML parsuje się i python-docx otwiera plik."""
    from xml.dom import minidom
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith('.xml'):
                minidom.parseString(z.read(n))
    try:
        import docx
        docx.Document(path)
    except ImportError:
        return 'OK (XML valid; python-docx niezainstalowany — pominięto)'
    except Exception as e:  # noqa: BLE001
        return 'UWAGA: python-docx nie otworzył pliku: %s' % e
    return 'OK'


def extract_text(path, name='word/document.xml'):
    """Zwraca czytelny tekst dokumentu (akapity, taby jako ' | ')."""
    import html
    with zipfile.ZipFile(path) as z:
        t = z.read(name).decode('utf-8')
    t = re.sub(r'</w:p>', '\n', t)
    t = re.sub(r'<w:tab[^>]*/>', ' | ', t)
    t = re.sub(r'<[^>]+>', '', t)
    return html.unescape(t)


def keyword_coverage(path, keywords):
    """Zwraca (pokryte, brakujące) po dopasowaniu bez uwzględniania wielkości liter.
    Uwaga: dla polskiej odmiany podawaj rdzenie/formy występujące w CV."""
    low = extract_text(path).lower()
    low = re.sub(r'\s+', ' ', low)
    missing = [k for k in keywords if k.lower() not in low]
    return [k for k in keywords if k not in missing], missing
